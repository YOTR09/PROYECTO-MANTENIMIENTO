import os
import tempfile
import unittest
import openpyxl
from config.database import init_db, get_db_cursor
from src.controllers.reporte_controller import ReporteController
from src.repositories.socio_repository import SocioRepository
from src.repositories.vehiculo_repository import VehiculoRepository
from src.repositories.mantenimiento_repository import MantenimientoRepository
from src.repositories.tipo_mantenimiento_repository import TipoMantenimientoRepository


class TestReportesFiltros(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        os.environ["TESTING"] = "1"
        init_db()

    def setUp(self):
        with get_db_cursor(commit=True) as cur:
            cur.execute("DELETE FROM historial_mantenimiento")
            cur.execute("DELETE FROM mantenimiento_programado")
            cur.execute("DELETE FROM vehiculo")
            cur.execute("DELETE FROM socio")

        self.s1 = SocioRepository.create("V-1", "Socio Uno", "0412-1", "Activo")
        self.s2 = SocioRepository.create("V-2", "Socio Dos", "0412-2", "Inactivo")
        self.s3 = SocioRepository.create("V-3", "Socio Tres Sin Unidades", "0412-3", "Activo")

        self.v1 = VehiculoRepository.create(self.s1, "01", "AAA111", "Encava", 2020, 50000, "Activo")
        self.v2 = VehiculoRepository.create(self.s1, "02", "BBB222", "Yutong", 2019, 200000, "En Taller")
        self.v3 = VehiculoRepository.create(self.s2, "03", "CCC333", "Iveco", 2018, 120000, "Activo")

        tipos = TipoMantenimientoRepository.get_all(solo_activos=True)
        self.t1, self.t2 = tipos[0]["id_tipo"], tipos[1]["id_tipo"]

        MantenimientoRepository.add_historial(self.v1, self.t1, "2026-01-10", 40000, 100.0, "Taller Central", "a")
        MantenimientoRepository.add_historial(self.v1, self.t2, "2026-02-15", 45000, 200.0, "Frenos Palmar", "b")
        MantenimientoRepository.add_historial(self.v2, self.t1, "2026-03-20", 190000, 300.0, "Taller Central", "c")

        # Programaciones: una vencida (km), una al día
        MantenimientoRepository.upsert_programacion(self.v2, self.t1, "2026-01-01", 190000, "2099-01-01", 195000)
        MantenimientoRepository.upsert_programacion(self.v1, self.t1, "2026-01-01", 40000, "2099-06-01", 90000)

    # ---------- Validación de filtros ----------
    def test_normalizacion_todos_a_none(self):
        f = ReporteController.normalizar_filtros({"status": "Todos", "id_socio": None, "taller": "  ", "fecha_desde": ""})
        self.assertTrue(all(v is None for v in f.values()))

    def test_fechas_incoherentes_o_invalidas(self):
        with self.assertRaises(Exception):
            ReporteController.normalizar_filtros({"fecha_desde": "2026-05-01", "fecha_hasta": "2026-01-01"})
        with self.assertRaises(Exception):
            ReporteController.normalizar_filtros({"fecha_desde": "01/05/2026"})
        with self.assertRaises(Exception):
            ReporteController.normalizar_filtros({"km_min": "500", "km_max": "100"})
        ok, _, msg = ReporteController.contar_registros(3, {"fecha_desde": "basura"})
        self.assertFalse(ok)
        self.assertIn("AAAA-MM-DD", msg)

    # ---------- Conteos por reporte ----------
    def test_flota_filtros(self):
        self.assertEqual(ReporteController.contar_registros(1, {})[1], 3)
        self.assertEqual(ReporteController.contar_registros(1, {"status": "En Taller"})[1], 1)
        self.assertEqual(ReporteController.contar_registros(1, {"id_socio": self.s1})[1], 2)
        self.assertEqual(ReporteController.contar_registros(1, {"km_min": 100000})[1], 2)
        self.assertEqual(ReporteController.contar_registros(1, {"km_min": 100000, "km_max": 150000})[1], 1)

    def test_preventivo_filtros(self):
        self.assertEqual(ReporteController.contar_registros(2, {})[1], 2)
        self.assertEqual(ReporteController.contar_registros(2, {"id_vehiculo": self.v2})[1], 1)
        # v2 (km actual 200000 > 195000) está vencido
        self.assertEqual(ReporteController.contar_registros(2, {"estado_alerta": "Vencido"})[1], 1)
        self.assertEqual(ReporteController.contar_registros(2, {"estado_alerta": "Críticos"})[1], 1)
        self.assertEqual(ReporteController.contar_registros(2, {"fecha_hasta": "2050-01-01"})[1], 0)

    def test_historial_filtros(self):
        self.assertEqual(ReporteController.contar_registros(3, {})[1], 3)
        self.assertEqual(ReporteController.contar_registros(3, {"fecha_desde": "2026-02-01", "fecha_hasta": "2026-02-28"})[1], 1)
        self.assertEqual(ReporteController.contar_registros(3, {"taller": "Taller Central"})[1], 2)
        self.assertEqual(ReporteController.contar_registros(3, {"id_vehiculo": self.v1, "id_tipo": self.t2})[1], 1)
        self.assertEqual(ReporteController.contar_registros(5, {"taller": "Inexistente"})[1], 0)

    def test_socios_filtros(self):
        self.assertEqual(ReporteController.contar_registros(4, {})[1], 3)
        self.assertEqual(ReporteController.contar_registros(4, {"estado": "Inactivo"})[1], 1)
        self.assertEqual(ReporteController.contar_registros(4, {"tenencia": "Con unidades"})[1], 2)
        self.assertEqual(ReporteController.contar_registros(4, {"tenencia": "Sin unidades"})[1], 1)
        self.assertEqual(ReporteController.contar_registros(4, {"id_socio": self.s1})[1], 1)

    def test_talleres_unicos(self):
        c = ReporteController.obtener_catalogos_para_filtros()
        self.assertEqual(c["talleres"], ["Frenos Palmar", "Taller Central"])

    # ---------- Generación de archivos ----------
    def test_generacion_con_filtros(self):
        with tempfile.TemporaryDirectory() as d:
            p1, p2, p3, p4, p5 = (os.path.join(d, n) for n in ("1.pdf", "2.pdf", "3.xlsx", "4.pdf", "5.xlsx"))
            self.assertTrue(ReporteController.generar_reporte_1_flota_pdf(p1, {"status": "En Taller"})[0])
            self.assertTrue(ReporteController.generar_reporte_2_preventivo_pdf(p2, {"estado_alerta": "Vencido"})[0])
            self.assertTrue(ReporteController.generar_reporte_3_historial_excel(p3, {"taller": "Taller Central"})[0])
            self.assertTrue(ReporteController.generar_reporte_4_socios_pdf(p4, {"estado": "Activo"})[0])
            self.assertTrue(ReporteController.generar_reporte_5_metricas_excel(p5, {"fecha_desde": "2026-01-01"})[0])
            for p in (p1, p2, p4):
                with open(p, "rb") as fh:
                    self.assertTrue(fh.read().startswith(b"%PDF"))

            ws = openpyxl.load_workbook(p3).active
            self.assertIn("Taller Central", ws["A3"].value)
            self.assertEqual(ws["H7"].value, "=SUM(H5:H6)")  # 2 filas filtradas

    def test_generacion_sin_resultados(self):
        with tempfile.TemporaryDirectory() as d:
            p1 = os.path.join(d, "1.pdf")
            p3 = os.path.join(d, "3.xlsx")
            p5 = os.path.join(d, "5.xlsx")
            p4 = os.path.join(d, "4.pdf")
            self.assertTrue(ReporteController.generar_reporte_1_flota_pdf(p1, {"km_min": 9999999})[0])
            self.assertTrue(ReporteController.generar_reporte_3_historial_excel(p3, {"taller": "Nada"})[0])
            self.assertTrue(ReporteController.generar_reporte_5_metricas_excel(p5, {"taller": "Nada"})[0])
            self.assertTrue(ReporteController.generar_reporte_4_socios_pdf(p4, {"estado": "Inactivo", "tenencia": "Sin unidades"})[0])
            with open(p1, "rb") as fh:
                self.assertTrue(fh.read().startswith(b"%PDF"))

    def test_generacion_con_filtros_invalidos(self):
        with tempfile.TemporaryDirectory() as d:
            ok, msg = ReporteController.generar_reporte_3_historial_excel(
                os.path.join(d, "x.xlsx"), {"fecha_desde": "2026-12-01", "fecha_hasta": "2026-01-01"}
            )
            self.assertFalse(ok)
            self.assertIn("posterior", msg)


if __name__ == "__main__":
    unittest.main()
