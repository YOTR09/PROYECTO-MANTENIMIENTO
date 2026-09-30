import os
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from typing import List, Dict, Any

class ExcelReportGenerator:

    @staticmethod
    def _aplicar_estilos_base(ws, titulo: str, subtitulo: str):
        # Encabezado institucional
        ws.merge_cells("A1:H1")
        cell_title = ws["A1"]
        cell_title.value = f"EMPRESA DE TRANSPORTE BRISAS DEL PALMAR — {titulo.upper()}"
        cell_title.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
        cell_title.fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        cell_title.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 30

        ws.merge_cells("A2:H2")
        cell_sub = ws["A2"]
        ahora_str = datetime.now().strftime("%d/%m/%Y %H:%M")
        cell_sub.value = f"{subtitulo} | Generado el {ahora_str}"
        cell_sub.font = Font(name="Calibri", size=10, italic=True, color="475569")
        cell_sub.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 20

    @staticmethod
    def _autoajustar_columnas(ws):
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                if len(val) > max_len and cell.row > 2:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # ---------------- REPORTE 3: BITÁCORA HISTÓRICA Y COSTOS ----------------
    @staticmethod
    def generar_reporte_historial_costos(historial: List[Dict[str, Any]], ruta_destino: str):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Historial y Costos"
        ws.views.sheetView[0].showGridLines = True

        ExcelReportGenerator._aplicar_estilos_base(
            ws,
            "Bitácora Histórica de Mantenimiento y Costos",
            "Relación de servicios mecánicos ejecutados y gastos asociados"
        )

        headers = ["ID", "Fecha", "Unidad", "Placa", "Rutina Ejecutada", "Odómetro al Servicio", "Taller / Mecánico", "Costo ($)", "Detalle del Trabajo"]
        ws.row_dimensions[4].height = 24

        font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        fill_header = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
        border_thin = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1")
        )

        for col_num, h in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col_num, value=h)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border_thin

        row_idx = 5
        fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

        for h in historial:
            ws.cell(row=row_idx, column=1, value=h.get("id_historial")).alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=2, value=h.get("fecha_realizado")).alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=3, value=h.get("numero_unidad")).alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=4, value=h.get("placa")).alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=5, value=h.get("tipo_nombre")).alignment = Alignment(horizontal="left")
            
            c_km = ws.cell(row=row_idx, column=6, value=h.get("km_al_momento", 0))
            c_km.number_format = "#,##0"
            c_km.alignment = Alignment(horizontal="right")

            ws.cell(row=row_idx, column=7, value=h.get("taller_mecanico") or "-").alignment = Alignment(horizontal="left")

            c_costo = ws.cell(row=row_idx, column=8, value=float(h.get("costo") or 0.0))
            c_costo.number_format = "$#,##0.00"
            c_costo.alignment = Alignment(horizontal="right")

            ws.cell(row=row_idx, column=9, value=h.get("descripcion_trabajo") or "-").alignment = Alignment(horizontal="left")

            # Bordes y zebra
            for col_num in range(1, 10):
                c = ws.cell(row=row_idx, column=col_num)
                c.border = border_thin
                if row_idx % 2 == 0:
                    c.fill = fill_zebra

            row_idx += 1

        # Fila de totales con fórmulas de Excel
        ws.row_dimensions[row_idx].height = 22
        cell_lbl_total = ws.cell(row=row_idx, column=7, value="TOTAL GASTO OPERATIVO:")
        cell_lbl_total.font = Font(name="Calibri", size=11, bold=True, color="1E3A8A")
        cell_lbl_total.alignment = Alignment(horizontal="right", vertical="center")

        cell_sum = ws.cell(row=row_idx, column=8)
        if len(historial) > 0:
            cell_sum.value = f"=SUM(H5:H{row_idx-1})"
        else:
            cell_sum.value = 0.0
        cell_sum.font = Font(name="Calibri", size=11, bold=True, color="1E3A8A")
        cell_sum.number_format = "$#,##0.00"
        cell_sum.alignment = Alignment(horizontal="right", vertical="center")
        cell_sum.fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
        cell_sum.border = Border(top=Side(style="double", color="1E3A8A"), bottom=Side(style="double", color="1E3A8A"))

        ExcelReportGenerator._autoajustar_columnas(ws)
        wb.save(ruta_destino)

    # ---------------- REPORTE 5: RESUMEN EJECUTIVO Y MÉTRICAS ----------------
    @staticmethod
    def generar_reporte_metricas_talleres(historial: List[Dict[str, Any]], vehiculos: List[Dict[str, Any]], ruta_destino: str):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Métricas Ejecutivas"
        ws.views.sheetView[0].showGridLines = True

        ExcelReportGenerator._aplicar_estilos_base(
            ws,
            "Resumen Ejecutivo de Talleres y Desempeño Operativo",
            "Métricas consolidadas de costos por proveedor y tipos de mantenimiento"
        )

        font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        fill_sec1 = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        fill_sec2 = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")
        border_thin = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1")
        )

        # SECCIÓN 1: Desglose por Taller Mecánico
        ws.cell(row=4, column=1, value="1. CONSOLIDADO DE GASTOS POR TALLER / PROVEEDOR").font = Font(bold=True, size=11, color="1E3A8A")
        headers1 = ["Taller / Tallerista", "Servicios Realizados", "Gasto Acumulado ($)", "Costo Promedio / Servicio ($)"]
        for c, h in enumerate(headers1, 1):
            cell = ws.cell(row=5, column=c, value=h)
            cell.font = font_header
            cell.fill = fill_sec1
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border_thin

        # Agrupar datos por taller
        talleres_data = {}
        for h in historial:
            taller = h.get("taller_mecanico") or "Taller no especificado / Interno"
            costo = float(h.get("costo") or 0.0)
            if taller not in talleres_data:
                talleres_data[taller] = {"count": 0, "total": 0.0}
            talleres_data[taller]["count"] += 1
            talleres_data[taller]["total"] += costo

        curr_row = 6
        for t_nom, t_info in sorted(talleres_data.items(), key=lambda x: x[1]["total"], reverse=True):
            ws.cell(row=curr_row, column=1, value=t_nom).alignment = Alignment(horizontal="left")
            c_cnt = ws.cell(row=curr_row, column=2, value=t_info["count"])
            c_cnt.number_format = "#,##0"
            c_cnt.alignment = Alignment(horizontal="center")
            
            c_tot = ws.cell(row=curr_row, column=3, value=t_info["total"])
            c_tot.number_format = "$#,##0.00"
            c_tot.alignment = Alignment(horizontal="right")

            prom = t_info["total"] / t_info["count"] if t_info["count"] > 0 else 0.0
            c_avg = ws.cell(row=curr_row, column=4, value=prom)
            c_avg.number_format = "$#,##0.00"
            c_avg.alignment = Alignment(horizontal="right")

            for col in range(1, 5):
                ws.cell(row=curr_row, column=col).border = border_thin
            curr_row += 1

        # SECCIÓN 2: Desglose por Rutina de Mantenimiento
        curr_row += 2
        ws.cell(row=curr_row, column=1, value="2. INTERVENCIONES POR TIPO DE RUTINA PREVENTIVA").font = Font(bold=True, size=11, color="1E3A8A")
        curr_row += 1

        headers2 = ["Rutina Preventiva", "Frecuencia de Ejecución", "Inversión Total ($)"]
        for c, h in enumerate(headers2, 1):
            cell = ws.cell(row=curr_row, column=c, value=h)
            cell.font = font_header
            cell.fill = fill_sec2
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border_thin

        rutinas_data = {}
        for h in historial:
            rutina = h.get("tipo_nombre") or "General"
            costo = float(h.get("costo") or 0.0)
            if rutina not in rutinas_data:
                rutinas_data[rutina] = {"count": 0, "total": 0.0}
            rutinas_data[rutina]["count"] += 1
            rutinas_data[rutina]["total"] += costo

        curr_row += 1
        for r_nom, r_info in sorted(rutinas_data.items(), key=lambda x: x[1]["count"], reverse=True):
            ws.cell(row=curr_row, column=1, value=r_nom).alignment = Alignment(horizontal="left")
            c_cnt = ws.cell(row=curr_row, column=2, value=r_info["count"])
            c_cnt.number_format = "#,##0"
            c_cnt.alignment = Alignment(horizontal="center")

            c_tot = ws.cell(row=curr_row, column=3, value=r_info["total"])
            c_tot.number_format = "$#,##0.00"
            c_tot.alignment = Alignment(horizontal="right")

            for col in range(1, 4):
                ws.cell(row=curr_row, column=col).border = border_thin
            curr_row += 1

        ExcelReportGenerator._autoajustar_columnas(ws)
        wb.save(ruta_destino)
